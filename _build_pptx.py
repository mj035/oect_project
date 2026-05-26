"""중간 발표용 PPT 빌더.

oect_midterm.pptx 생성. 16:9 와이드, 다크 테마, 이미지 placeholder 포함.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ── 컬러 (사이트 다크 테마와 매칭) ────────────────────────────
BG       = RGBColor(0x0F, 0x14, 0x19)   # 어두운 배경
SURFACE  = RGBColor(0x1A, 0x21, 0x2E)
SURFACE2 = RGBColor(0x24, 0x2C, 0x3A)
TEXT     = RGBColor(0xE6, 0xEC, 0xF1)
MUTED    = RGBColor(0x9C, 0xA3, 0xAF)
ACCENT   = RGBColor(0x60, 0xA5, 0xFA)   # 파랑
GREEN    = RGBColor(0x34, 0xD3, 0x99)
ORANGE   = RGBColor(0xFB, 0xBF, 0x24)
RED      = RGBColor(0xF8, 0x71, 0x71)
LINE     = RGBColor(0x37, 0x41, 0x51)

# ── 슬라이드 사이즈 (16:9) ─────────────────────────────────
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height
BLANK = prs.slide_layouts[6]


def add_bg(slide, color=BG):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg


def add_text(slide, text, left, top, width, height, *,
             size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="맑은 고딕"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor

    lines = text.split("\n") if isinstance(text, str) else text
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_chip(slide, text, left, top, color=ACCENT, width=None):
    """라운드 사각형 칩 (뱃지). 텍스트 길이 기반 자동 폭."""
    if width is None:
        width = Inches(0.4 + 0.13 * len(text))
    height = Inches(0.42)
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    sh.adjustments[0] = 0.5
    sh.fill.solid(); sh.fill.fore_color.rgb = SURFACE2
    sh.line.color.rgb = color
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = Emu(50000); tf.margin_right = Emu(50000)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = "맑은 고딕"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = color
    return sh


def add_image_placeholder(slide, left, top, width, height, caption=""):
    """이미지 자리 placeholder (점선 박스)."""
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid(); box.fill.fore_color.rgb = SURFACE
    box.line.color.rgb = MUTED
    box.line.width = Pt(1)
    box.line.dash_style = 7  # dash
    box.shadow.inherit = False
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(100000); tf.margin_right = Emu(100000)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "🖼  " + (caption or "이미지 자리")
    r.font.name = "맑은 고딕"
    r.font.size = Pt(14)
    r.font.color.rgb = MUTED
    return box


def add_card(slide, left, top, width, height, accent=ACCENT):
    """둥근 카드 컨테이너."""
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    c.adjustments[0] = 0.04
    c.fill.solid(); c.fill.fore_color.rgb = SURFACE
    c.line.color.rgb = accent
    c.line.width = Pt(1.5)
    c.shadow.inherit = False
    c.text_frame.text = ""
    return c


def add_divider_line(slide, top, color=LINE):
    line = slide.shapes.add_connector(1, Inches(0.7), top, SW - Inches(0.7), top)
    line.line.color.rgb = color
    line.line.width = Pt(1)
    return line


def add_page_header(slide, idx, total, section, title):
    add_text(slide, f"{idx:02d} / {total:02d}",
             Inches(0.7), Inches(0.35), Inches(1.5), Inches(0.4),
             size=12, color=MUTED)
    add_text(slide, section,
             Inches(2.2), Inches(0.35), Inches(6), Inches(0.4),
             size=12, bold=True, color=ACCENT)
    add_text(slide, title,
             Inches(0.7), Inches(0.85), Inches(12), Inches(0.8),
             size=30, bold=True, color=TEXT)
    add_divider_line(slide, Inches(1.75))


# ════════════════════════════════════════════════════════════
# Slide 01 — 표지
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)

# 좌측 강조 바
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.3), SH)
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT
bar.line.fill.background(); bar.shadow.inherit = False

add_text(s, "AI 활용 프로젝트 — 중간 발표",
         Inches(0.9), Inches(0.7), Inches(11), Inches(0.5),
         size=14, color=MUTED)

add_text(s, "OECT 논문 검색 & AI 연구 도우미",
         Inches(0.9), Inches(1.5), Inches(12), Inches(1.6),
         size=44, bold=True, color=TEXT)

add_text(s, "유기 전기화학 트랜지스터(OECT) 논문을\n모션 인식과 AI로 탐색하고 분석하는 시스템",
         Inches(0.9), Inches(3.0), Inches(11), Inches(1.4),
         size=18, color=MUTED)

# 팀 정보 카드
add_card(s, Inches(0.9), Inches(5.0), Inches(11.5), Inches(1.9), accent=ACCENT)
add_text(s, "TEAM",
         Inches(1.2), Inches(5.15), Inches(2), Inches(0.4),
         size=12, bold=True, color=ACCENT)
members = ["김동우  ( 학번:                 )",
           "조혜원  ( 학번:                 )",
           "손윤성  ( 학번:                 )",
           "오민재  ( 학번:                 )"]
for i, m in enumerate(members):
    col = i % 2
    row = i // 2
    add_text(s, m,
             Inches(1.2 + col * 5.5), Inches(5.65 + row * 0.55), Inches(5), Inches(0.45),
             size=15, color=TEXT)

add_text(s, "2026 학부 발표 · 제출일 ___ / ___",
         Inches(0.9), Inches(7.05), Inches(11), Inches(0.3),
         size=10, color=MUTED, align=PP_ALIGN.RIGHT)


# ════════════════════════════════════════════════════════════
# Slide 02 — 프로젝트 개요
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 2, 15, "01. PROJECT OVERVIEW", "프로젝트 개요")

add_text(s,
         "유기 전기화학 트랜지스터(OECT) 분야의 논문을\n"
         "모션 인식으로 탐색하고, AI로 통합 분석하며,\n"
         "회로까지 자동으로 시각화하는 웹 기반 연구 도우미.",
         Inches(0.7), Inches(2.0), Inches(12), Inches(2.0),
         size=22, color=TEXT)

# 3-카드 핵심 요약
labels = [
    ("모션 입력",     "카메라 + 손 제스처로\n원하는 장기/기능 선택", GREEN),
    ("AI 통합 분석",  "논문 PDF에서 6개 항목을\n구조화된 형태로 추출",  ACCENT),
    ("회로 시각화",   "논문에서 회로 추출 + AI가\n부족한 부분 자동 보완",  ORANGE),
]
card_w = Inches(3.9); card_h = Inches(2.3); gap = Inches(0.15)
total_w = card_w * 3 + gap * 2
start_x = (SW - total_w) / 2
for i, (title, desc, col) in enumerate(labels):
    left = start_x + (card_w + gap) * i
    add_card(s, left, Inches(4.5), card_w, card_h, accent=col)
    add_text(s, title,
             left + Inches(0.3), Inches(4.7), card_w - Inches(0.6), Inches(0.5),
             size=18, bold=True, color=col)
    add_text(s, desc,
             left + Inches(0.3), Inches(5.3), card_w - Inches(0.6), Inches(1.4),
             size=13, color=TEXT)


# ════════════════════════════════════════════════════════════
# Slide 03 — 왜 이 프로젝트인가
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 3, 15, "02. MOTIVATION", "왜 이 프로젝트인가")

add_text(s, "ChatGPT만으로는 안 되는가?",
         Inches(0.7), Inches(2.0), Inches(12), Inches(0.6),
         size=22, bold=True, color=ORANGE)

add_text(s,
         "• 단순 “논문 요약” 수준이면 ChatGPT 한 번 호출로 끝난다\n"
         "• 그러나 “OECT 분야 특화 + 실제 연구 워크플로우 통합”은 다르다",
         Inches(0.7), Inches(2.7), Inches(12), Inches(1.3),
         size=15, color=TEXT)

# 차별화 포인트 4개
diff = [
    ("도메인 특화",       "OECT 6필드 구조화 분석 — 일반 LLM은 못 함"),
    ("멀티모달 입력",     "모션 인식으로 카테고리 진입 — 발표 임팩트"),
    ("회로 자동 추출",    "논문 figure 인식 + 부족하면 AI 자동 생성"),
    ("정직한 신뢰도 표시", "AI 결과의 출처/신뢰도를 사용자에게 노출"),
]
for i, (h, d) in enumerate(diff):
    row = i // 2
    col = i % 2
    left = Inches(0.7 + col * 6.1)
    top  = Inches(4.4 + row * 1.2)
    add_card(s, left, top, Inches(5.9), Inches(1.05), accent=ACCENT)
    add_text(s, h,
             left + Inches(0.25), top + Inches(0.1), Inches(5.5), Inches(0.4),
             size=15, bold=True, color=ACCENT)
    add_text(s, d,
             left + Inches(0.25), top + Inches(0.55), Inches(5.5), Inches(0.5),
             size=12, color=TEXT)


# ════════════════════════════════════════════════════════════
# Slide 04 — 시스템 한눈에
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 4, 15, "03. ARCHITECTURE", "시스템 한눈에")

# 가로 흐름 5단계
steps = [
    ("USER",         "카메라 +\n제스처",        MUTED),
    ("MOTION",       "MediaPipe\nHolistic",     GREEN),
    ("SEARCH",       "OpenAlex\n논문 API",      ACCENT),
    ("AI ANALYSIS",  "Gemini API\n6필드 분석",   ORANGE),
    ("OUTPUT",       "회로 / 재료 /\n공정 / 응용",   RED),
]
n = len(steps)
total_w = Inches(11.5)
step_w  = total_w / n
start_x = (SW - total_w) / 2
y = Inches(2.6)

for i, (label, desc, col) in enumerate(steps):
    left = start_x + step_w * i
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                             left + Emu(50000), y, step_w - Emu(100000), Inches(2.0))
    box.adjustments[0] = 0.1
    box.fill.solid(); box.fill.fore_color.rgb = SURFACE
    box.line.color.rgb = col; box.line.width = Pt(2)
    box.shadow.inherit = False
    box.text_frame.text = ""
    add_text(s, label,
             left + Emu(50000), y + Inches(0.2), step_w - Emu(100000), Inches(0.45),
             size=13, bold=True, color=col, align=PP_ALIGN.CENTER)
    add_text(s, desc,
             left + Emu(50000), y + Inches(0.8), step_w - Emu(100000), Inches(1.0),
             size=12, color=TEXT, align=PP_ALIGN.CENTER)
    if i < n - 1:
        # 화살표
        ax = left + step_w - Emu(60000)
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                   ax, y + Inches(0.85),
                                   Emu(120000), Inches(0.3))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = MUTED
        arrow.line.fill.background()
        arrow.shadow.inherit = False

# 하단 보강 설명
add_text(s,
         "└ 기술 스택 :  Flask · Python 3.12 · PyMuPDF · schemdraw · MediaPipe · Gemini · OpenAlex",
         Inches(0.7), Inches(5.5), Inches(12), Inches(0.5),
         size=12, color=MUTED, align=PP_ALIGN.CENTER)
add_text(s,
         "└ 협업 환경 :  Claude Code (코드 어시스트) · Git · VS Code",
         Inches(0.7), Inches(6.0), Inches(12), Inches(0.5),
         size=12, color=MUTED, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# Slide 05 — 기술① 모션 인식
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 5, 15, "04. TECHNOLOGY ①", "모션 인식 — 카메라로 장기 선택")

add_text(s, "MediaPipe Holistic + 제스처 매칭",
         Inches(0.7), Inches(2.0), Inches(7), Inches(0.5),
         size=18, bold=True, color=GREEN)

gestures = [
    ("손 → 머리 위",       "뇌 (Brain)"),
    ("손 → 가슴",          "심장 (Cardiac)"),
    ("손 → 배",            "간 (Liver)"),
    ("V사인 + 얼굴 옆",    "눈 (Eye)"),
    ("양손 주먹",          "근육 (Muscle)"),
    ("양손 손바닥 펼침",   "피부 (Skin)"),
]
for i, (g, organ) in enumerate(gestures):
    add_text(s, f"•  {g}   →   {organ}",
             Inches(0.7), Inches(2.7 + i * 0.42), Inches(7), Inches(0.4),
             size=14, color=TEXT)

add_text(s, "1.5초 유지 시 자동 페이지 이동",
         Inches(0.7), Inches(5.5), Inches(7), Inches(0.4),
         size=13, color=ORANGE)

# 우측 캡처 자리
add_image_placeholder(s, Inches(8.0), Inches(2.0),
                      Inches(4.7), Inches(4.5),
                      "메인 화면 (모션 인식) 캡처")


# ════════════════════════════════════════════════════════════
# Slide 06 — 기술② AI 분석 6필드
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 6, 15, "04. TECHNOLOGY ②", "AI 분석 — 6개 항목 구조화 추출")

add_text(s, "Gemini API (gemini-2.5-flash) + PDF 첨부 분석",
         Inches(0.7), Inches(2.0), Inches(8), Inches(0.5),
         size=16, bold=True, color=ACCENT)

fields = [
    ("재료",        "사용된 모든 폴리머/기판/용매 추출",         GREEN),
    ("디바이스 구조","레이어 구성, 채널 지오메트리, 전극",       ACCENT),
    ("동작 원리",    "센싱/동작 메커니즘 (이온 결합 등)",          ORANGE),
    ("공정",        "제작 단계 순서대로 정리",                    GREEN),
    ("응용/타겟",   "타겟 분석물 + 성능 지표 (검출한계 등)",     ACCENT),
    ("회로",        "별도 추출 — 다음 슬라이드에서 자세히",       ORANGE),
]
for i, (name, desc, col) in enumerate(fields):
    row = i // 2; col_idx = i % 2
    left = Inches(0.7 + col_idx * 4.1); top = Inches(2.7 + row * 1.18)
    add_card(s, left, top, Inches(3.9), Inches(1.0), accent=col)
    add_text(s, name,
             left + Inches(0.2), top + Inches(0.08), Inches(3.5), Inches(0.4),
             size=14, bold=True, color=col)
    add_text(s, desc,
             left + Inches(0.2), top + Inches(0.5), Inches(3.5), Inches(0.5),
             size=11, color=TEXT)

add_image_placeholder(s, Inches(9.0), Inches(2.7),
                      Inches(3.7), Inches(4.0),
                      "분석 결과 6탭 화면 캡처")


# ════════════════════════════════════════════════════════════
# Slide 07 — 기술③ 회로 추출 (자랑 포인트)
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 7, 15, "04. TECHNOLOGY ③", "회로 추출 — 3계층 폴백 ★")

# 3 단계 박스
tiers = [
    ("1순위",  "raster 임베디드 이미지\n(PNG/JPEG로 박힌 회로)",  GREEN,  "PyMuPDF"),
    ("2순위",  "vector 회로 영역 추출\n(페이지 렌더 + Gemini ROI)", ACCENT, "PyMuPDF + Gemini"),
    ("3순위",  "텍스트 기반 자동 생성\n(schemdraw 템플릿)",         ORANGE, "Gemini + schemdraw"),
]
for i, (k, d, col, tech) in enumerate(tiers):
    left = Inches(0.7 + i * 4.2); top = Inches(2.5)
    add_card(s, left, top, Inches(4.0), Inches(2.6), accent=col)
    add_text(s, k,
             left + Inches(0.3), top + Inches(0.15), Inches(3.5), Inches(0.5),
             size=18, bold=True, color=col)
    add_text(s, d,
             left + Inches(0.3), top + Inches(0.85), Inches(3.5), Inches(1.2),
             size=14, color=TEXT)
    add_text(s, f"⚙  {tech}",
             left + Inches(0.3), top + Inches(2.0), Inches(3.5), Inches(0.4),
             size=11, color=MUTED)
    if i < 2:
        ax = left + Inches(4.0)
        arrow = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                   ax + Emu(20000), top + Inches(1.1),
                                   Emu(160000), Inches(0.4))
        arrow.fill.solid(); arrow.fill.fore_color.rgb = MUTED
        arrow.line.fill.background(); arrow.shadow.inherit = False

# 하단 캡처
add_image_placeholder(s, Inches(0.7), Inches(5.5),
                      Inches(5.8), Inches(1.7),
                      "회로 탭 — raster / 벡터 추출 / AI 생성 비교")
add_image_placeholder(s, Inches(6.7), Inches(5.5),
                      Inches(6.0), Inches(1.7),
                      "AI 생성 회로 + 신뢰도 게이지")


# ════════════════════════════════════════════════════════════
# Slide 08 — 기술④ 개발 환경 / 협업 도구
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 8, 15, "04. TECHNOLOGY ④", "개발 환경과 협업 도구")

tools = [
    ("Claude Code", "AI 코드 어시스트 — 빠른 프로토타이핑 / 디버깅 / 리팩토링",            ACCENT),
    ("Gemini API",  "PDF·이미지 멀티모달 분석 + 카테고리 분류 + 회로 토폴로지 추출",       ORANGE),
    ("OpenAlex",    "OECT 분야 논문 메타데이터 + OA PDF 검색 (무료, API 키 불필요)",        GREEN),
    ("MediaPipe",   "브라우저 카메라로 손/얼굴 랜드마크 실시간 검출 (CDN 클라이언트)",     GREEN),
    ("PyMuPDF",     "PDF 텍스트/이미지 추출 + 페이지 렌더링 (벡터 회로 폴백)",              ACCENT),
    ("schemdraw",   "Python에서 회로도 자동 생성 — OECT 전용 심볼 직접 정의",                ORANGE),
    ("Flask",       "Python 백엔드 웹 프레임워크 — REST API 엔드포인트 7개",                 ACCENT),
    ("Git + VSCode","팀 협업 + 코드 관리 + 통합 개발 환경",                                  GREEN),
]
for i, (name, desc, col) in enumerate(tools):
    row = i // 2; col_idx = i % 2
    left = Inches(0.7 + col_idx * 6.0)
    top = Inches(2.1 + row * 1.13)
    add_card(s, left, top, Inches(5.9), Inches(1.0), accent=col)
    add_text(s, name,
             left + Inches(0.2), top + Inches(0.08), Inches(2), Inches(0.4),
             size=14, bold=True, color=col)
    add_text(s, desc,
             left + Inches(0.2), top + Inches(0.5), Inches(5.6), Inches(0.5),
             size=11, color=TEXT)


# ════════════════════════════════════════════════════════════
# Slide 09 — 어려움 1: 회로 그리기
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 9, 15, "05. CHALLENGE ①", "회로를 어떻게 보여줄까?")

add_text(s, "처음 시도 — 실패한 접근들",
         Inches(0.7), Inches(2.0), Inches(11), Inches(0.5),
         size=20, bold=True, color=RED)

failed = [
    ("× LLM이 schemdraw 코드 직접 생성",       "→ 매번 깨진 회로, topology 커버리지 부족"),
    ("× LLM이 SVG 직접 그리기",                  "→ 좌표가 어긋나고 라벨 충돌"),
    ("× 텍스트 기반 mermaid 다이어그램",         "→ 회로 기호(트랜지스터 등) 표현 불가"),
]
for i, (h, d) in enumerate(failed):
    add_text(s, h,
             Inches(0.7), Inches(2.8 + i * 0.95), Inches(11), Inches(0.5),
             size=16, bold=True, color=RED)
    add_text(s, d,
             Inches(1.2), Inches(3.25 + i * 0.95), Inches(11), Inches(0.5),
             size=13, color=MUTED)

add_card(s, Inches(0.7), Inches(5.85), Inches(12), Inches(1.3), accent=ORANGE)
add_text(s, "💡  결국 깨달음",
         Inches(0.95), Inches(5.95), Inches(11), Inches(0.5),
         size=14, bold=True, color=ORANGE)
add_text(s,
         "AI에게 회로를 \"창작\"하게 시키지 말자.\n"
         "AI는 분류·추출에 능하고, 그림은 우리가 만든 템플릿이 그리게 하자.",
         Inches(0.95), Inches(6.4), Inches(11.5), Inches(0.7),
         size=13, color=TEXT)


# ════════════════════════════════════════════════════════════
# Slide 10 — 어려움 2: PDF 차단 / Quota
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 10, 15, "05. CHALLENGE ②", "외부 의존성의 한계")

c1_left = Inches(0.7); c2_left = Inches(6.95)
top = Inches(2.0); h = Inches(4.8)

# 카드 1: PDF 차단
add_card(s, c1_left, top, Inches(6.0), h, accent=RED)
add_text(s, "출판사 PDF 자동 다운로드 차단",
         c1_left + Inches(0.3), top + Inches(0.2), Inches(5.4), Inches(0.5),
         size=16, bold=True, color=RED)
add_text(s,
         "• Wiley / Elsevier / ACS 등은 봇 방지 차단\n"
         "• User-Agent 위장으로도 우회 어려움\n"
         "• 자동화의 가장 큰 현실적 장벽",
         c1_left + Inches(0.3), top + Inches(0.85), Inches(5.4), Inches(1.7),
         size=12, color=TEXT)
add_text(s, "📌 임시 해결",
         c1_left + Inches(0.3), top + Inches(2.7), Inches(5.4), Inches(0.4),
         size=13, bold=True, color=ORANGE)
add_text(s,
         "사용자가 PDF를 직접 다운로드 후 업로드 →\n"
         "/api/analyze-upload 엔드포인트로 처리",
         c1_left + Inches(0.3), top + Inches(3.15), Inches(5.4), Inches(1.3),
         size=11, color=TEXT)

# 카드 2: Quota
add_card(s, c2_left, top, Inches(6.0), h, accent=RED)
add_text(s, "Gemini API 무료 티어 한계",
         c2_left + Inches(0.3), top + Inches(0.2), Inches(5.4), Inches(0.5),
         size=16, bold=True, color=RED)
add_text(s,
         "• 분당 10회 / 일일 250회 한도\n"
         "• 논문 1편 분석 = 5~6회 호출\n"
         "• 발표 시연 도중 막히면 데모 사고",
         c2_left + Inches(0.3), top + Inches(0.85), Inches(5.4), Inches(1.7),
         size=12, color=TEXT)
add_text(s, "📌 임시 해결",
         c2_left + Inches(0.3), top + Inches(2.7), Inches(5.4), Inches(0.4),
         size=13, bold=True, color=ORANGE)
add_text(s,
         "여러 Gemini 모델 자동 폴백 +\n"
         "재시도 백오프(3 → 6 → 10 → 15초)",
         c2_left + Inches(0.3), top + Inches(3.15), Inches(5.4), Inches(1.3),
         size=11, color=TEXT)


# ════════════════════════════════════════════════════════════
# Slide 11 — 해결책 1: 3계층 폴백
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 11, 15, "06. SOLUTION ①", "3계층 회로 폴백 — 빈 결과 0%")

add_text(s, "어떤 논문이든 회로 1장은 무조건 표시",
         Inches(0.7), Inches(2.0), Inches(12), Inches(0.5),
         size=18, bold=True, color=GREEN)

# 흐름도 (수직)
flow = [
    ("PDF 입력",                                  None,    MUTED),
    ("1순위 │ raster 회로 추출",                  GREEN,   "→ 있으면 즉시 표시 ✓"),
    ("2순위 │ vector 영역 검출 + 크롭",            ACCENT,  "→ 페이지 렌더 후 Gemini가 좌표 추출"),
    ("3순위 │ 텍스트 → AI 토폴로지 분류 → 템플릿", ORANGE,  "→ 어떤 논문이든 1장 보장"),
]
y = Inches(2.7); step = Inches(1.0)
for i, (label, col, note) in enumerate(flow):
    if col is None:
        add_text(s, "▶  " + label,
                 Inches(0.7), y + step * i, Inches(11), Inches(0.5),
                 size=15, bold=True, color=MUTED)
    else:
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(0.9), y + step * i,
                                 Inches(7.4), Inches(0.7))
        box.adjustments[0] = 0.3
        box.fill.solid(); box.fill.fore_color.rgb = SURFACE
        box.line.color.rgb = col; box.line.width = Pt(2)
        box.shadow.inherit = False; box.text_frame.text = ""
        add_text(s, label,
                 Inches(1.1), y + step * i + Inches(0.15), Inches(7), Inches(0.5),
                 size=14, bold=True, color=col)
        if note:
            add_text(s, note,
                     Inches(8.5), y + step * i + Inches(0.2), Inches(4.5), Inches(0.5),
                     size=11, color=TEXT)

# 하단 캡션
add_text(s,
         "│  각 회로마다 출처 메타 데이터 부착  →  사용자가 “원본 vs AI 생성”을 한눈에 구분",
         Inches(0.7), Inches(6.6), Inches(12), Inches(0.5),
         size=12, color=MUTED)


# ════════════════════════════════════════════════════════════
# Slide 12 — 해결책 2: 출처 뱃지 + 신뢰도 + 캐싱
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 12, 15, "06. SOLUTION ②", "정직한 표시 + 발표 안정성")

add_text(s, "AI 결과의 \"신뢰도\"를 사용자에게 노출",
         Inches(0.7), Inches(2.0), Inches(12), Inches(0.5),
         size=18, bold=True, color=ACCENT)

# 좌측: 뱃지 / 신뢰도
add_card(s, Inches(0.7), Inches(2.7), Inches(7), Inches(2.5), accent=ACCENT)
add_text(s, "회로 출처 뱃지",
         Inches(0.95), Inches(2.85), Inches(6), Inches(0.5),
         size=14, bold=True, color=ACCENT)
add_chip(s, "논문 원본",   Inches(0.95), Inches(3.45), color=GREEN)
add_chip(s, "벡터 추출",   Inches(2.55), Inches(3.45), color=ACCENT)
add_chip(s, "AI 생성",     Inches(4.15), Inches(3.45), color=ORANGE)
add_text(s,
         "+ AI 생성 회로엔 신뢰도 게이지 바\n"
         "+ 캡션에 토폴로지 + 신뢰도 명시\n"
         "+ 도메인 룰로 confidence 후처리 보정",
         Inches(0.95), Inches(4.05), Inches(6.6), Inches(1.1),
         size=12, color=TEXT)

# 우측: 캐싱
add_card(s, Inches(8.0), Inches(2.7), Inches(4.7), Inches(2.5), accent=GREEN)
add_text(s, "분석 결과 캐싱",
         Inches(8.25), Inches(2.85), Inches(4.5), Inches(0.5),
         size=14, bold=True, color=GREEN)
add_text(s,
         "• PDF SHA-256으로 캐시 키\n"
         "• 같은 논문 두 번째부터 즉시 응답\n"
         "• 발표용 논문 사전 캐싱 →\n   본 발표 도중 quota 사고 0%",
         Inches(8.25), Inches(3.45), Inches(4.5), Inches(1.7),
         size=11, color=TEXT)

# 하단 캡처
add_image_placeholder(s, Inches(0.7), Inches(5.5),
                      Inches(12.0), Inches(1.7),
                      "회로 카드 — 뱃지 + 신뢰도 게이지 UI 캡처")


# ════════════════════════════════════════════════════════════
# Slide 13 — 향후 계획 (최종본)
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 13, 15, "07. NEXT STEPS", "최종본까지의 계획")

phases = [
    ("Phase 1 │ 분석 정밀도 ↑",
     "재료 8카테고리 분할 (substrate · source · drain · channel · gate · electrolyte · interconnect · encapsulation) 도입",
     ACCENT),
    ("Phase 2 │ 검색 품질 ↑",
     "OpenAlex 다중 OA 미러 자동 시도 + stretchable / hydrogel / ionic gel 키워드 보강 → 차단률 50% ↓",
     GREEN),
    ("Phase 3 │ 통합 분석",
     "선택된 N편 논문을 합쳐 \"공통 재료 / 공통 공정 / 트렌드\" 통합 보고서 생성 (ChatGPT로 못 하는 차별화)",
     ORANGE),
    ("Phase 4 │ 발표 안정화",
     "시연용 5~10편 사전 캐시 + 데모 모드 + 백업 API 키 — 본 발표 사고 0% 보장",
     RED),
]
for i, (h, d, col) in enumerate(phases):
    top = Inches(2.1 + i * 1.18)
    add_card(s, Inches(0.7), top, Inches(12), Inches(1.05), accent=col)
    add_text(s, h,
             Inches(0.95), top + Inches(0.1), Inches(11), Inches(0.45),
             size=14, bold=True, color=col)
    add_text(s, d,
             Inches(0.95), top + Inches(0.55), Inches(11.5), Inches(0.5),
             size=11, color=TEXT)


# ════════════════════════════════════════════════════════════
# Slide 14 — 확장 방향
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)
add_page_header(s, 14, 15, "08. EXPANSION", "이 프로젝트의 확장 방향")

add_text(s, "OECT를 넘어선 일반화",
         Inches(0.7), Inches(2.0), Inches(12), Inches(0.5),
         size=20, bold=True, color=ACCENT)

ext = [
    ("타 분야 적용",
     "이차전지 / 광촉매 / 페로브스카이트 등 다른 재료공학 분야로 도메인 룰만 교체하면 즉시 확장 가능",
     GREEN),
    ("학술 연구 도우미 일반화",
     "검색 → 분석 → 시각화의 워크플로우는 분야 무관하게 적용 가능. 학부생 연구실 입문 도구로 발전",
     ACCENT),
    ("음성 + 모션 멀티모달",
     "한국어 음성 입력으로 행동 선택 + 모션으로 카테고리 → 키보드 없는 연구 인터페이스",
     ORANGE),
    ("논문 비교 리포트",
     "여러 논문을 자동 비교 표로 정리 + 1차 / 2차 실험 추천 조합 제시",
     GREEN),
    ("실험 노트 연동",
     "분석 결과를 학생 실험 노트(Notion / Markdown)로 자동 export + 인용 정리",
     ACCENT),
]
for i, (h, d, col) in enumerate(ext):
    top = Inches(2.65 + i * 0.85)
    add_card(s, Inches(0.7), top, Inches(12), Inches(0.75), accent=col)
    add_text(s, h,
             Inches(0.95), top + Inches(0.05), Inches(4), Inches(0.5),
             size=13, bold=True, color=col)
    add_text(s, d,
             Inches(4.95), top + Inches(0.05), Inches(7.5), Inches(0.7),
             size=11, color=TEXT, anchor=MSO_ANCHOR.MIDDLE)


# ════════════════════════════════════════════════════════════
# Slide 15 — 마무리 / 슬로건
# ════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK); add_bg(s)

# 좌측 강조 바
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.3), SH)
bar.fill.solid(); bar.fill.fore_color.rgb = ORANGE
bar.line.fill.background(); bar.shadow.inherit = False

add_text(s, "WRAP UP",
         Inches(0.9), Inches(0.7), Inches(11), Inches(0.5),
         size=14, color=MUTED)

add_text(s, "AI를 잘 쓴다는 건,\nAI의 한계를 잘 안다는 것이다.",
         Inches(0.9), Inches(1.8), Inches(12), Inches(2.2),
         size=42, bold=True, color=TEXT)

add_text(s,
         "\"잘하는 일을 시키고, 못하는 건 우리가 채운다\"",
         Inches(0.9), Inches(4.3), Inches(12), Inches(0.6),
         size=18, color=ORANGE)

add_text(s,
         "이 프로젝트는 AI에 모든 것을 맡기지 않고,\n"
         "도메인 지식 + 폴백 설계 + 정직한 결과 표시로\n"
         "“실제로 작동하는” AI 도구를 만드는 시도였다.",
         Inches(0.9), Inches(5.2), Inches(12), Inches(1.5),
         size=14, color=MUTED)

add_text(s, "Q & A",
         Inches(0.9), Inches(6.85), Inches(12), Inches(0.5),
         size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# 저장
# ════════════════════════════════════════════════════════════
out_path = r"C:/Users/dhals/oect-paper-search/oect_midterm.pptx"
prs.save(out_path)
print(f"OK saved: {out_path}")
print(f"slides: {len(prs.slides)}")
